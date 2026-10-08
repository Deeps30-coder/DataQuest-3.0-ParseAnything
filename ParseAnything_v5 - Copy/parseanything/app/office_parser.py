from .models import Block
from .pdf_parser import blocks_to_markdown


def parse_docx(path):
    """Reads a Word file in the same order it appears: paragraphs, tables and pictures."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    blocks = []
    n = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.split('}')[-1]

        if tag == 'p':
            n += 1
            para = Paragraph(child, doc)
            pictures = len(child.findall('.//' + qn('pic:pic')))
            if pictures:
                blocks.append(Block(id=f'p1_f{n}', type='figure', page=1, text='Picture in document',
                                    confidence=0.9, source='python-docx', metadata={'image_count': pictures}))
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or '').lower() if para.style else ''
            typ = 'heading' if 'heading' in style else ('list' if 'list' in style else 'paragraph')
            blocks.append(Block(id=f'p1_b{n}', type=typ, page=1, text=text, confidence=0.98,
                                source='python-docx', metadata={'style': style}))

        elif tag == 'tbl':
            n += 1
            table = Table(child, doc)
            rows = [[c.text.strip() for c in r.cells] for r in table.rows]
            blocks.append(Block(id=f'p1_t{n}', type='table', page=1, text='', confidence=0.96,
                                source='python-docx', metadata={'rows': rows}))

    return 1, blocks, blocks_to_markdown(blocks), [], 'success'


def parse_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    blocks = []
    for si, ws in enumerate(wb.worksheets, 1):
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(['' if v is None else str(v) for v in row])
        while rows and not any(rows[-1]):
            rows.pop()
        if rows:
            blocks.append(Block(id=f's{si}_t1', type='spreadsheet', page=si, text='', confidence=0.99,
                                source='openpyxl', metadata={'sheet': ws.title, 'rows': rows}))
    md = blocks_to_markdown(blocks)
    return max(1, len(wb.worksheets)), blocks, md, [], 'success'


def parse_pptx(path):
    """Reads a PowerPoint file: text, tables, and pictures or charts as figures."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    blocks = []
    for si, slide in enumerate(prs.slides, 1):
        for j, shape in enumerate(slide.shapes, 1):
            if getattr(shape, 'has_table', False) and shape.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                blocks.append(Block(id=f's{si}_t{j}', type='table', page=si, text='', confidence=0.96,
                                    source='python-pptx', metadata={'rows': rows, 'slide': si}))
                continue

            is_chart = bool(getattr(shape, 'has_chart', False))
            try:
                is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            except Exception:
                is_picture = False
            if is_chart or is_picture:
                kind = 'chart' if is_chart else 'picture'
                blocks.append(Block(id=f's{si}_f{j}', type='figure', page=si, text=f'{kind.capitalize()} on slide',
                                    confidence=0.9, source='python-pptx', metadata={'slide': si, 'kind': kind}))
                continue

            if not getattr(shape, 'has_text_frame', False):
                continue
            text = shape.text.strip()
            if not text:
                continue
            typ = 'heading' if j == 1 else 'slide'
            blocks.append(Block(id=f's{si}_b{j}', type=typ, page=si, text=text, confidence=0.97,
                                source='python-pptx', metadata={'slide': si}))

    return len(prs.slides), blocks, blocks_to_markdown(blocks), [], 'success'
